import os
import re
import logging
import subprocess
from typing import List, Dict, Any, Tuple
from pptx import Presentation
from sqlalchemy import select
from app.db.models import SlideChunk, SyllabusUnit, Subject

logger = logging.getLogger("ppt_parser")

UNIT_PATTERNS = [
    r"unit\s*[-_:]*\s*([1-5|i{1-3}|iv|v]+)",
    r"chapter\s*[-_:]*\s*([1-5|i{1-3}|iv|v]+)",
    r"module\s*[-_:]*\s*([1-5|i{1-3}|iv|v]+)"
]

ROMAN_MAP = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5}

def parse_unit_number(text: str) -> int:
    text_lower = text.lower()
    for pattern in UNIT_PATTERNS:
        match = re.search(pattern, text_lower)
        if match:
            val = match.group(1).strip()
            if val.isdigit():
                num = int(val)
                if 1 <= num <= 5:
                    return num
            elif val in ROMAN_MAP:
                return ROMAN_MAP[val]
    return 0

class PPTParser:
    """
    Parses PPT/PPTX files, extracts slide text, tables, and notes, detects units,
    and returns chunk records for pgvector indexing.
    """

    @staticmethod
    def convert_ppt_to_pptx(ppt_path: str, output_dir: str) -> str:
        """Converts legacy .ppt file to .pptx using LibreOffice CLI."""
        os.makedirs(output_dir, exist_ok=True)
        cmd = ["libreoffice", "--headless", "--convert-to", "pptx", ppt_path, "--outdir", output_dir]
        try:
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)
            filename = os.path.splitext(os.path.basename(ppt_path))[0] + ".pptx"
            converted_path = os.path.join(output_dir, filename)
            if os.path.exists(converted_path):
                return converted_path
            raise FileNotFoundError(f"LibreOffice conversion failed to produce {converted_path}")
        except Exception as e:
            logger.error(f"Error converting PPT {ppt_path}: {e}")
            raise RuntimeError(f"PPT conversion failed: {e}")

    @classmethod
    def parse_file(cls, file_path: str) -> List[Dict[str, Any]]:
        """
        Parses PPTX file (converting PPT if needed) and returns slide chunks.
        """
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".ppt":
            pptx_path = cls.convert_ppt_to_pptx(file_path, os.path.dirname(file_path))
        elif ext == ".pptx":
            pptx_path = file_path
        else:
            raise ValueError(f"Unsupported presentation file type: {ext}")

        prs = Presentation(pptx_path)
        filename = os.path.basename(file_path)
        filename_unit = parse_unit_number(filename)

        current_unit = filename_unit if filename_unit > 0 else 1
        chunks = []

        total_slides = len(prs.slides)
        for i, slide in enumerate(prs.slides):
            slide_no = i + 1
            slide_texts = []
            slide_title = ""

            # Extract shapes
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        if not slide_title:
                            slide_title = text.split("\n")[0]
                        slide_texts.append(text)

                if shape.has_table:
                    table_texts = []
                    for row in shape.table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_cells:
                            table_texts.append(" | ".join(row_cells))
                    if table_texts:
                        slide_texts.append("\nTable:\n" + "\n".join(table_texts))

            # Extract notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"Notes: {notes}")

            # Unit detection from slide title
            title_unit = parse_unit_number(slide_title) if slide_title else 0
            if title_unit > 0:
                current_unit = title_unit
            elif filename_unit == 0:
                # Distribute slides linearly across 5 units if no explicit title/filename units found
                current_unit = min(5, max(1, int(((slide_no - 1) / max(1, total_slides)) * 5) + 1))

            chunk_text = "\n\n".join(slide_texts).strip()
            if not chunk_text:
                chunk_text = f"Slide {slide_no} content for unit {current_unit}"

            chunks.append({
                "unit_no": current_unit,
                "ppt_name": filename,
                "slide_no": slide_no,
                "title": slide_title,
                "chunk_text": chunk_text
            })

        return chunks

    @classmethod
    async def process_and_store_ppts(
        cls,
        user_id: str,
        subject_id: str,
        ppt_paths: List[str],
        session
    ) -> List[SlideChunk]:
        """
        Parses all PPT files and stores SlideChunk records in database. Also auto-populates SyllabusUnits if missing.
        """
        stored_chunks = []
        unit_map: Dict[int, List[str]] = {u: [] for u in range(1, 6)}

        for path in ppt_paths:
            if not os.path.exists(path):
                continue
            chunks_data = cls.parse_file(path)
            for cd in chunks_data:
                unit_no = cd["unit_no"]
                chunk = SlideChunk(
                    user_id=user_id,
                    subject_id=subject_id,
                    unit_no=unit_no,
                    ppt_name=cd["ppt_name"],
                    slide_no=cd["slide_no"],
                    chunk_text=cd["chunk_text"],
                    metadata_json={"title": cd.get("title", "")}
                )
                session.add(chunk)
                stored_chunks.append(chunk)

                if cd.get("title"):
                    unit_map[unit_no].append(cd["title"])

        # Auto-create SyllabusUnit records if they don't exist
        for u_no in range(1, 6):
            res = await session.execute(
                select(SyllabusUnit).where(
                    SyllabusUnit.subject_id == subject_id,
                    SyllabusUnit.unit_no == u_no
                )
            )
            existing_unit = res.scalar_one_or_none()
            if not existing_unit:
                titles = unit_map[u_no]
                unit_title = titles[0] if titles else f"Unit {u_no}"
                unit = SyllabusUnit(
                    subject_id=subject_id,
                    unit_no=u_no,
                    title=unit_title,
                    topics=list(set(titles[:5]))
                )
                session.add(unit)

        await session.flush()
        return stored_chunks
